# Copyright 2026 EPAM Systems, Inc. (“EPAM”)
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
Tests for conversation pagination endpoints.
"""

import io
import shutil
from datetime import datetime, timezone
from unittest.mock import patch, MagicMock, AsyncMock

import pypandoc
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pypdf import PdfReader

import codemie.rest_api.routers.conversation as conversation_router
from codemie.chains.base import Thought
from codemie.rest_api.main import app
from codemie.rest_api.models.conversation import (
    Conversation,
    ConversationHistoryPaginationData,
    ConversationListItem,
    GeneratedMessage,
)
from codemie.rest_api.models.index import SortOrder
from codemie.rest_api.security.user import User


@pytest.fixture
def user():
    return User(id="user-123", username="testuser", name="Test User")


@pytest.fixture
def mock_litellm_startup():
    """Mock LiteLLM startup functions to prevent budget initialization errors."""
    with patch("codemie.rest_api.main.is_litellm_enabled", return_value=False):
        with patch("codemie.rest_api.main._initialize_database_and_defaults", return_value=None):
            with patch("codemie.rest_api.main.close_llm_proxy_client", new_callable=AsyncMock):
                yield


@pytest.fixture
def client(mock_litellm_startup):
    """Create TestClient with proper mocking."""
    return TestClient(app)


@pytest.fixture(autouse=True)
def override_dependency(user):
    app.dependency_overrides[conversation_router.authenticate] = lambda: user
    yield
    app.dependency_overrides = {}


@patch("codemie.rest_api.routers.conversation.Conversation.get_user_conversations", new_callable=MagicMock)
def test_get_conversations_without_pagination(mock_get, client):
    """
    Test GET /v1/conversations without pagination parameters.
    Should call Conversation.get_user_conversations.
    """
    t_first = datetime(2025, 1, 10, tzinfo=timezone.utc)
    t_last = datetime(2025, 1, 15, tzinfo=timezone.utc)
    mock_get.return_value = [
        ConversationListItem(
            id="conv-1",
            name="Conv 1",
            date=t_last,
            very_first_msg_at=t_first,
            very_last_msg_at=t_last,
        )
    ]

    response = client.get("/v1/conversations")

    assert response.status_code == 200
    data = response.json()
    assert data[0]["very_first_msg_at"] is not None
    assert data[0]["very_last_msg_at"] is not None
    mock_get.assert_called_once()


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_user_conversations_paginated", new_callable=MagicMock
)
def test_get_conversations_with_pagination(mock_paginated, client):
    """
    Test GET /v1/conversations with pagination parameters.
    Should call ConversationService.get_user_conversations_paginated.
    """
    t_first = datetime(2025, 1, 12, tzinfo=timezone.utc)
    t_last = datetime(2025, 1, 14, tzinfo=timezone.utc)
    mock_paginated.return_value = [
        ConversationListItem(
            id="conv-2",
            name="Conv 2",
            date=t_last,
            very_first_msg_at=t_first,
            very_last_msg_at=t_last,
        )
    ]

    response = client.get("/v1/conversations?page=0&per_page=10")

    assert response.status_code == 200
    data = response.json()
    assert data[0]["very_first_msg_at"] is not None
    assert data[0]["very_last_msg_at"] is not None
    mock_paginated.assert_called_once()


@patch("codemie.rest_api.routers.conversation.Conversation.get_user_conversations", new_callable=MagicMock)
def test_get_conversations_without_pagination_forwards_is_finished_filter(mock_get, client):
    mock_get.return_value = []

    response = client.get("/v1/conversations?isFinished=false")

    assert response.status_code == 200
    mock_get.assert_called_once_with(user_id="user-123", filters={"is_finished": False})


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_user_conversations_paginated", new_callable=MagicMock
)
def test_get_conversations_with_pagination_forwards_is_finished_filter(mock_paginated, client):
    mock_paginated.return_value = []

    response = client.get("/v1/conversations?page=0&per_page=10&isFinished=true")

    assert response.status_code == 200
    mock_paginated.assert_called_once_with(user_id="user-123", page=0, per_page=10, is_finished=True)


@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_json_without_pagination(mock_find, _mock_can, _mock_assistants, client):
    """
    Test GET /v1/conversations/{id}/export without pagination (JSON).
    Should call Conversation.find_by_id.
    Timestamps appear in the export when messages have dates; absent when they don't.
    """
    t_first = datetime(2025, 3, 1, 10, 0, tzinfo=timezone.utc)
    t_last = datetime(2025, 3, 1, 11, 0, tzinfo=timezone.utc)
    mock_find.return_value = Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[
            GeneratedMessage(role="User", message="Hello", date=t_first),
            GeneratedMessage(role="Assistant", message="Hi", date=t_last),
        ],
    )

    response = client.get("/v1/conversations/conv-123/export?export_format=json")

    assert response.status_code == 200
    data = response.json()
    assert data["very_first_msg_at"] is not None
    assert data["very_last_msg_at"] is not None
    mock_find.assert_called_once_with("conv-123")


@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_json_without_pagination_no_dates(mock_find, _mock_can, _mock_assistants, client):
    """
    Timestamps are absent from the export JSON when no messages have a date value.
    """
    mock_find.return_value = Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[GeneratedMessage(role="User", message="Hello")],
    )

    response = client.get("/v1/conversations/conv-123/export?export_format=json")

    assert response.status_code == 200
    data = response.json()
    assert "very_first_msg_at" not in data
    assert "very_last_msg_at" not in data
    mock_find.assert_called_once_with("conv-123")


@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
def test_export_json_with_pagination(mock_slice, _mock_can, _mock_assistants, client):
    """
    Test GET /v1/conversations/{id}/export with pagination (JSON).
    Should call ConversationService.get_conversation_history_slice.
    Timestamps from the service (full-conversation bounds) appear in the export.
    """
    t_first = datetime(2025, 3, 1, 9, 0, tzinfo=timezone.utc)
    t_last = datetime(2025, 3, 1, 11, 0, tzinfo=timezone.utc)
    conv = Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[GeneratedMessage(role="User", message="Hello", date=t_first)],
    )
    mock_slice.return_value = (conv, 1, t_first, t_last)

    response = client.get("/v1/conversations/conv-123/export?export_format=json&page=0&per_page=50")

    assert response.status_code == 200
    data = response.json()
    assert data["very_first_msg_at"] is not None
    assert data["very_last_msg_at"] is not None
    mock_slice.assert_called_once()


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice",
    new_callable=MagicMock,
    return_value=(None, 0, None, None),
)
def test_export_json_not_found(mock_slice, client):
    """
    Test GET /v1/conversations/{id}/export with pagination when not found.
    Should return 404.
    """
    response = client.get("/v1/conversations/non-existent/export?export_format=json&page=0&per_page=10")

    assert response.status_code == 404


def _conversation_with_thought_and_empty_pair():
    return Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[
            GeneratedMessage(role="User", message="Explain X", history_index=0),
            GeneratedMessage(
                role="Assistant",
                message="Here is the explanation.",
                history_index=0,
                thoughts=[Thought(id="t1", message="ran search tool", author_name="Search Tool", author_type="Tool")],
            ),
            GeneratedMessage(role="User", message="Run a tool silently", history_index=1),
            GeneratedMessage(
                role="Assistant",
                message="",
                history_index=1,
                thoughts=[Thought(id="t2", message="tool output only", author_name="Tool", author_type="Tool")],
            ),
        ],
    )


def _conversation_with_empty_pair_and_unexportable_thought():
    """Second pair: the assistant produced no text, and its only thought node is one
    ExportUtils.should_include_thought filters out anyway (blank message). Concise mode strips
    nothing from this pair - full mode would not render that thought either - so the pair must
    survive in both modes rather than taking the empty-pair skip."""
    return Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[
            GeneratedMessage(role="User", message="Explain X", history_index=0),
            GeneratedMessage(role="Assistant", message="Here is the explanation.", history_index=0),
            GeneratedMessage(role="User", message="Did the tool run?", history_index=1),
            GeneratedMessage(
                role="Assistant",
                message="",
                history_index=1,
                thoughts=[Thought(id="t3", message="   ", author_name="Tool", author_type="Tool")],
            ),
        ],
    )


def _conversation_with_code_block():
    return Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[
            GeneratedMessage(role="User", message="Show me code", history_index=0),
            GeneratedMessage(
                role="Assistant",
                message="Here is code:\n\n```python\nprint('hello world')\n```\n",
                history_index=0,
            ),
        ],
    )


def _conversation_with_list_table_link():
    return Conversation(
        id="conv-123",
        conversation_id="conv-123",
        user_id="user-123",
        history=[
            GeneratedMessage(role="User", message="Summarize with formatting", history_index=0),
            GeneratedMessage(
                role="Assistant",
                message=(
                    "Summary:\n\n"
                    "- Item one\n"
                    "- Item two\n\n"
                    "| Col A | Col B |\n"
                    "| --- | --- |\n"
                    "| a1 | b1 |\n\n"
                    "See [our site](https://example.com) for more.\n"
                ),
                history_index=0,
            ),
        ],
    )


def _write_source_as_output(*, source, to, outputfile, **_kwargs):
    """Stand-in for pypandoc.convert_text: writes the serialized markdown Pandoc would have
    received directly to the output file, so tests can assert on the real DocumentBuilder/
    MessageExporter pipeline output without depending on the pandoc/pdflatex binaries."""
    with open(outputfile, "w", encoding="utf-8") as f:
        f.write(source)


# --- Wiring checks (parameter reaches MessageExporter with the right value) ---


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@pytest.mark.parametrize("query_suffix", ["", "&include_tool_outputs=false"])
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
@patch("codemie.rest_api.routers.conversation.MessageExporter")
def test_export_concise_mode_omits_thoughts(
    mock_exporter_cls, mock_find, _mock_can, _mock_assistants, client, export_format, query_suffix
):
    """Omitted or explicit include_tool_outputs=false must construct MessageExporter with
    include_tool_outputs=False for both pdf and docx."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()
    mock_instance = MagicMock()
    mock_instance.run.return_value = iter([b"bytes"])
    mock_instance.content_type = "application/octet-stream"
    mock_exporter_cls.return_value = mock_instance

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}{query_suffix}")

    assert response.status_code == 200
    mock_exporter_cls.assert_called_once()
    assert mock_exporter_cls.call_args.kwargs["include_tool_outputs"] is False


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
@patch("codemie.rest_api.routers.conversation.MessageExporter")
def test_export_full_mode_keeps_thoughts(
    mock_exporter_cls, mock_find, _mock_can, _mock_assistants, client, export_format
):
    """include_tool_outputs=true must construct MessageExporter with include_tool_outputs=True."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()
    mock_instance = MagicMock()
    mock_instance.run.return_value = iter([b"bytes"])
    mock_instance.content_type = "application/octet-stream"
    mock_exporter_cls.return_value = mock_instance

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}&include_tool_outputs=true")

    assert response.status_code == 200
    assert mock_exporter_cls.call_args.kwargs["include_tool_outputs"] is True


@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_json_unaffected_by_include_tool_outputs_param(mock_find, _mock_can, _mock_assistants, client):
    """json export must not error and must ignore include_tool_outputs when the frontend sends it anyway."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()

    response = client.get("/v1/conversations/conv-123/export?export_format=json&include_tool_outputs=true")

    assert response.status_code == 200


# --- Non-mocked, real-pipeline checks (closes the zero-coverage gap the AC calls out) ---
# MessageExporter/DocumentBuilder run for real; only pypandoc.convert_text's external binary
# call is stubbed, so assertions are against the actual serialized document content.


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_concise_mode_omits_thought_content_end_to_end(
    mock_find, _mock_can, _mock_assistants, _mock_convert, client, export_format
):
    """Real pipeline, thought content absent and the thought-only empty pair skipped entirely
    in concise mode (parameter omitted)."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}")

    assert response.status_code == 200
    text = response.content.decode("utf-8")
    assert "Search Tool" not in text
    assert "ran search tool" not in text
    assert "Run a tool silently" not in text  # empty pair's user message must not appear: pair fully skipped
    assert "Here is the explanation." in text  # surviving pair's assistant message still present


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_concise_mode_explicit_false_omits_thought_content_end_to_end(
    mock_find, _mock_can, _mock_assistants, _mock_convert, client, export_format
):
    """CR-004: real pipeline with include_tool_outputs=false passed explicitly on the query
    string must produce output identical (thought-absent, empty pair skipped) to the
    omitted-default case covered by test_export_concise_mode_omits_thought_content_end_to_end -
    guards against a default-handling divergence between 'omitted' and 'explicit false'."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}&include_tool_outputs=false")

    assert response.status_code == 200
    text = response.content.decode("utf-8")
    assert "Search Tool" not in text
    assert "ran search tool" not in text
    assert "Run a tool silently" not in text  # empty pair's user message must not appear: pair fully skipped
    assert "Here is the explanation." in text  # surviving pair's assistant message still present


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_full_mode_keeps_thought_content_end_to_end(
    mock_find, _mock_can, _mock_assistants, _mock_convert, client, export_format
):
    """Real pipeline, include_tool_outputs=true reproduces today's full export: thought content
    present and the empty-message/thought-only pair NOT skipped (regression: full variant
    unaffected)."""
    mock_find.return_value = _conversation_with_thought_and_empty_pair()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}&include_tool_outputs=true")

    assert response.status_code == 200
    text = response.content.decode("utf-8")
    assert "Search Tool" in text
    assert "Run a tool silently" in text  # empty pair's user heading/message still emitted, not skipped


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@pytest.mark.parametrize("query_suffix", ["", "&include_tool_outputs=true"])
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_keeps_empty_pair_whose_thoughts_were_never_exportable(
    mock_find, _mock_can, _mock_assistants, _mock_convert, client, export_format, query_suffix
):
    """The empty-pair skip must key on concise mode having actually removed something. A pair
    whose thoughts are all filtered out by should_include_thought renders identically in both
    modes, so concise mode must not drop the user's question - that content is not tool output."""
    mock_find.return_value = _conversation_with_empty_pair_and_unexportable_thought()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}{query_suffix}")

    assert response.status_code == 200
    text = response.content.decode("utf-8")
    assert "Did the tool run?" in text
    assert "No content to export in concise mode." not in text


# --- Real rendering checks (no pypandoc mock - genuine external binary conversion) ---
# These need binaries the pipeline shells out to, so each one is guarded. A missing binary must
# skip rather than fail: the behaviour under test is the document content, not the environment.


def _pandoc_available() -> bool:
    """pyproject depends on pypandoc-binary, which ships the pandoc executable inside the venv -
    so this is normally true even where `which pandoc` finds nothing on PATH. It goes false only
    on a platform with no pypandoc-binary wheel, where the install falls back to the sdist."""
    try:
        pypandoc.get_pandoc_path()
    except OSError:
        return False
    return True


requires_pandoc = pytest.mark.skipif(
    not _pandoc_available(),
    reason="pandoc binary unavailable (no pypandoc-binary wheel for this platform)",
)

# Only the PDF branch needs a LaTeX engine: PDFProcessor.get_pandoc_args passes
# --pdf-engine=pdflatex, and pandoc alone cannot produce a PDF without it.
requires_pdflatex = pytest.mark.skipif(
    shutil.which("pdflatex") is None,
    reason="pdflatex unavailable; pandoc cannot render PDF without a LaTeX engine",
)


def _is_code_style(paragraph) -> bool:
    """True for a paragraph carrying pandoc's DOCX code-block style. Pandoc's reference doc calls
    it 'Source Code', but the exact spelling python-docx reports has varied across pandoc versions
    ('SourceCode'), so compare case- and space-insensitively rather than pinning one spelling."""
    style = paragraph.style
    name = getattr(style, "name", None) if style is not None else None
    return bool(name) and name.replace(" ", "").lower() == "sourcecode"


@requires_pandoc
@pytest.mark.parametrize("query_suffix", ["", "&include_tool_outputs=true"])
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_docx_renders_fenced_code_block_as_code_run(
    mock_find, _mock_can, _mock_assistants, client, query_suffix, tmp_path
):
    """CR-003: full real pandoc pipeline (pypandoc.convert_text not mocked) for DOCX - a fenced
    code block in the surviving message renders as a genuine DOCX code-block run, in both
    concise mode (parameter omitted) and full mode (include_tool_outputs=true)."""
    mock_find.return_value = _conversation_with_code_block()

    response = client.get(f"/v1/conversations/conv-123/export?export_format=docx{query_suffix}")

    assert response.status_code == 200
    docx_path = tmp_path / "export.docx"
    docx_path.write_bytes(response.content)
    doc = DocxDocument(str(docx_path))

    code_paragraphs = [p for p in doc.paragraphs if _is_code_style(p)]
    assert any("print" in p.text for p in code_paragraphs)


@requires_pandoc
@pytest.mark.parametrize("query_suffix", ["", "&include_tool_outputs=true"])
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_docx_renders_list_table_link_correctly(
    mock_find, _mock_can, _mock_assistants, client, query_suffix, tmp_path
):
    """CR-005: full real pandoc pipeline for DOCX - a list, a table, and a link in the surviving
    message continue to render correctly in both concise mode (parameter omitted) and full mode
    (include_tool_outputs=true); formatting fidelity must not depend on which branch produced it."""
    mock_find.return_value = _conversation_with_list_table_link()

    response = client.get(f"/v1/conversations/conv-123/export?export_format=docx{query_suffix}")

    assert response.status_code == 200
    docx_path = tmp_path / "export.docx"
    docx_path.write_bytes(response.content)
    doc = DocxDocument(str(docx_path))

    paragraph_texts = [p.text for p in doc.paragraphs]
    assert "Item one" in paragraph_texts
    assert "Item two" in paragraph_texts

    assert len(doc.tables) == 1
    table_cells = [cell.text for row in doc.tables[0].rows for cell in row.cells]
    assert "Col A" in table_cells
    assert "a1" in table_cells

    hyperlink_targets = [rel.target_ref for rel in doc.part.rels.values() if "hyperlink" in rel.reltype]
    assert "https://example.com" in hyperlink_targets


@requires_pandoc
@requires_pdflatex
@pytest.mark.parametrize("query_suffix", ["", "&include_tool_outputs=true"])
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_pdf_renders_real_pdf_document(mock_find, _mock_can, _mock_assistants, client, query_suffix):
    """CR-005: the DOCX branch had real-render coverage and the PDF branch had none - every other
    PDF assertion stops at the markdown handed to pandoc. This is the only test that proves the
    PDF branch's own pandoc args (--pdf-engine=pdflatex, --listings, the fvextra header-includes)
    actually produce a readable PDF rather than a LaTeX error, in both modes."""
    mock_find.return_value = _conversation_with_list_table_link()

    response = client.get(f"/v1/conversations/conv-123/export?export_format=pdf{query_suffix}")

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")

    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "Item one" in text
    assert "Item two" in text
    assert "Col A" in text


@pytest.mark.parametrize(
    "export_format,expected_args,absent_args",
    [
        ("pdf", ["--pdf-engine=pdflatex", "--listings"], []),
        # DOCXProcessor.get_pandoc_args returns [], so the DOCX call must carry none of the
        # LaTeX-only flags - a processor mix-up would make DOCX export fail at the binary.
        ("docx", [], ["--pdf-engine=pdflatex", "--listings"]),
    ],
)
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_passes_format_specific_pandoc_args(
    mock_find, _mock_can, _mock_assistants, mock_convert, client, export_format, expected_args, absent_args
):
    """CR-005: runnable everywhere (the binary call is stubbed) counterpart to the pdflatex-guarded
    render test above - asserts the export target and the format-specific pandoc arguments the
    real render depends on, so the PDF branch's wiring stays covered on a host without LaTeX."""
    mock_find.return_value = _conversation_with_code_block()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}")

    assert response.status_code == 200
    kwargs = mock_convert.call_args.kwargs
    assert kwargs["to"] == export_format
    extra_args = kwargs["extra_args"]
    for arg in expected_args:
        assert arg in extra_args
    for arg in absent_args:
        assert arg not in extra_args
    if export_format == "pdf":
        # fvextra is what keeps long code lines from overflowing the page; it arrives as a
        # '-V header-includes=...' value rather than a flag of its own.
        assert any("fvextra" in arg for arg in extra_args)


@pytest.mark.parametrize("export_format", ["pdf", "docx"])
@patch("codemie.service.conversation.message_exporter.pypandoc.convert_text", side_effect=_write_source_as_output)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_export_list_table_link_content_survives_concise_filtering_end_to_end(
    mock_find, _mock_can, _mock_assistants, _mock_convert, client, export_format
):
    """CR-005: real DocumentBuilder/MessageExporter pipeline for both pdf and docx - list, table,
    and link markdown reach the pandoc input intact after passing through the concise-mode
    gating logic (only pypandoc.convert_text's external binary call is stubbed)."""
    mock_find.return_value = _conversation_with_list_table_link()

    response = client.get(f"/v1/conversations/conv-123/export?export_format={export_format}")

    assert response.status_code == 200
    text = response.content.decode("utf-8")
    assert "- Item one" in text
    assert "- Item two" in text
    assert "| Col A | Col B |" in text
    assert "[our site](https://example.com)" in text


# ---------------------------------------------------------------------------
# Helpers for GET /v1/conversations/{conversation_id} tests
# ---------------------------------------------------------------------------


def _make_test_conversation(n_messages: int = 3) -> Conversation:
    """Build a test Conversation with n_messages history items."""
    t_base = datetime(2025, 6, 1, tzinfo=timezone.utc)
    history = [
        GeneratedMessage(
            role="User" if i % 2 == 0 else "Assistant",
            message=f"msg-{i}",
            date=t_base.replace(hour=i),
        )
        for i in range(n_messages)
    ]
    return Conversation(
        id="conv-abc",
        conversation_id="conv-abc",
        conversation_name="Test Conversation",
        user_id="user-123",
        history=history,
        is_workflow_conversation=False,
        assistant_ids=["asst-1"],
    )


def _make_pagination(page=0, per_page=2, total=5, pages=3, has_next=True, has_previous=False):
    return ConversationHistoryPaginationData(
        page=page,
        per_page=per_page,
        total=total,
        pages=pages,
        has_next=has_next,
        has_previous=has_previous,
    )


# ---------------------------------------------------------------------------
# Backward compatibility
# ---------------------------------------------------------------------------


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
@patch("codemie.rest_api.routers.conversation.Conversation.find_by_id", new_callable=MagicMock)
def test_get_conversation_no_params_pagination_absent(mock_find, _mock_can, _mock_assistants, mock_slice, client):
    """
    When no pagination/sort params are provided, the response must not include
    a 'pagination' key and must return the full history unchanged.
    """
    mock_find.return_value = _make_test_conversation(n_messages=3)

    response = client.get("/v1/conversations/conv-abc")

    assert response.status_code == 200
    data = response.json()
    assert "pagination" not in data
    assert len(data["history"]) == 3
    mock_slice.assert_not_called()


# ---------------------------------------------------------------------------
# Pagination params present
# ---------------------------------------------------------------------------


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
def test_get_conversation_with_page_and_per_page(_mock_can, _mock_assistants, mock_slice, client):
    """
    When page and per_page are provided, get_conversation_history_slice is called and
    the response includes a correct pagination block.
    """
    conv = _make_test_conversation(n_messages=5)
    sliced_conv = _make_test_conversation(n_messages=2)
    sliced_conv.history = conv.history[:2]
    sliced_conv.id = conv.id
    sliced_conv.conversation_id = conv.conversation_id
    mock_slice.return_value = (sliced_conv, 5, None, None)

    response = client.get("/v1/conversations/conv-abc?page=0&per_page=2")

    assert response.status_code == 200
    data = response.json()
    assert "pagination" in data
    assert data["pagination"]["page"] == 0
    assert data["pagination"]["per_page"] == 2
    assert data["pagination"]["total"] == 5
    assert data["pagination"]["pages"] == 3
    assert data["pagination"]["has_next"] is True
    assert data["pagination"]["has_previous"] is False
    assert len(data["history"]) == 2
    mock_slice.assert_called_once_with(
        conversation_id="conv-abc",
        page=0,
        per_page=2,
        sort_order=None,
    )


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
def test_get_conversation_with_only_sort_order(_mock_can, _mock_assistants, mock_slice, client):
    """
    When only sort_order is provided (no page/per_page), it still triggers pagination
    using default page/per_page values, and the response includes a pagination block.
    """
    conv = _make_test_conversation(n_messages=3)
    mock_slice.return_value = (conv, 3, None, None)

    response = client.get("/v1/conversations/conv-abc?sort_order=asc")

    assert response.status_code == 200
    data = response.json()
    assert "pagination" in data
    mock_slice.assert_called_once()
    _, call_kwargs = mock_slice.call_args
    assert call_kwargs["sort_order"] == SortOrder.ASC


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
@patch("codemie.rest_api.routers.conversation.Assistant.get_by_ids", new_callable=MagicMock, return_value=[])
@patch("codemie.rest_api.routers.conversation.Ability.can", new_callable=MagicMock, return_value=True)
def test_get_conversation_sort_order_desc(_mock_can, _mock_assistants, mock_slice, client):
    """
    DESC sort_order is forwarded to get_conversation_history_slice correctly.
    """
    conv = _make_test_conversation(n_messages=3)
    sliced_conv = _make_test_conversation(n_messages=2)
    sliced_conv.history = conv.history[:2]
    sliced_conv.id = conv.id
    sliced_conv.conversation_id = conv.conversation_id
    mock_slice.return_value = (sliced_conv, 3, None, None)

    response = client.get("/v1/conversations/conv-abc?sort_order=desc&page=0&per_page=2")

    assert response.status_code == 200
    mock_slice.assert_called_once()
    _, call_kwargs = mock_slice.call_args
    assert call_kwargs["sort_order"] == SortOrder.DESC


# ---------------------------------------------------------------------------
# Error cases
# ---------------------------------------------------------------------------


@patch(
    "codemie.rest_api.routers.conversation.ConversationService.get_conversation_history_slice", new_callable=MagicMock
)
def test_get_conversation_not_found(mock_slice, client):
    """404 is returned when conversation does not exist, even with pagination params."""
    mock_slice.return_value = (None, 0, None, None)

    response = client.get("/v1/conversations/no-such-id?page=0&per_page=10")

    assert response.status_code == 404


def test_get_conversation_invalid_sort_order(client):
    """422 is returned when sort_order has an invalid value."""
    response = client.get("/v1/conversations/conv-abc?sort_order=invalid_value")

    assert response.status_code == 422


def test_get_conversation_negative_page(client):
    """422 is returned when page is negative (ge=0 constraint)."""
    response = client.get("/v1/conversations/conv-abc?page=-1")

    assert response.status_code == 422


def test_get_conversation_zero_per_page(client):
    """422 is returned when per_page is 0 (ge=1 constraint)."""
    response = client.get("/v1/conversations/conv-abc?per_page=0")

    assert response.status_code == 422
